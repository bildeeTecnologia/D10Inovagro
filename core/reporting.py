from io import BytesIO
from decimal import Decimal, ROUND_HALF_UP

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from .models import FieldEvent


def money(value):
    return f"R$ {value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _event_label(event_type, rubric_name=""):
    if event_type == "rubric":
        return rubric_name or "Evento de rubrica configurada"
    return dict(FieldEvent.EVENT_CHOICES).get(event_type, event_type)


def build_report_data(queryset, report_type):
    if report_type == "monthly":
        rows = [{"competencia": f"{r.competency.month:02d}/{r.competency.year}", "colaborador": r.collaborator.name, "unidade": r.unit.name, "dias_normais": r.normal_days, "dias_chuva": r.rain_days, "normal": r.normal_amount, "chuva": r.rain_amount, "total": r.total_amount} for r in queryset.select_related("competency", "collaborator", "unit")]
        return rows, []
    if report_type == "events":
        rows = [{"competencia": f"{e.competency.month:02d}/{e.competency.year}", "unidade": e.unit.name, "colaborador": e.collaborator.name, "tipo": e.get_event_type_display(), "data": e.event_date.strftime("%d/%m/%Y"), "quantidade": float(e.quantity), "status": e.get_status_display(), "motivo": e.reason} for e in queryset.select_related("unit", "collaborator", "competency")]
        inconsistencies = [r for r in rows if r["status"] in {"Pendente", "Rejeitado"}]
        return rows, inconsistencies
    if report_type == "event_costs":
        status_labels = dict(FieldEvent.STATUS_CHOICES)
        rows = []
        for item in queryset:
            total = item["total"] or Decimal("0.00")
            event_count = item["event_count"] or 0
            average = (total / Decimal(event_count)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP) if event_count else Decimal("0.00")
            rows.append({
                "competencia": f"{item['competency__month']:02d}/{item['competency__year']}",
                "unidade": f"{item['event__unit__name']} — {item['event__unit__company__legal_name']} (CNPJ {item['event__unit__company__cnpj']})",
                "evento": _event_label(item["event__event_type"], item.get("event__application__rubric__name") or ""),
                "status": status_labels.get(item["event__status"], item["event__status"]),
                "quantidade": event_count,
                "total": total,
                "media": average,
            })
        return rows, []
    rows = [{"competencia": f"{r['competency__month']:02d}/{r['competency__year']}", "unidade": r["event__unit__name"], "total": r["total"] or Decimal("0.00")} for r in queryset]
    return rows, []


def create_xlsx(report_type, rows, inconsistencies, metadata):
    wb = Workbook()
    summary = wb.active
    summary.title = "Resumo"
    summary.append(["Relatório", metadata["title"]])
    summary.append(["Emissor", metadata["user"]])
    summary.append(["Emissão", metadata["issued_at"]])
    summary.append(["Filtros", metadata["filters_text"]])
    summary.append(["Total de linhas", len(rows)])
    summary.append([])
    detail = wb.create_sheet("Detalhamento")
    if rows:
        headers = list(rows[0].keys())
    elif report_type == "cost":
        headers = ["competencia", "unidade", "total"]
    elif report_type == "event_costs":
        headers = ["competencia", "unidade", "evento", "status", "quantidade", "total", "media"]
    elif report_type == "monthly":
        headers = ["competencia", "colaborador", "unidade", "dias_normais", "dias_chuva", "normal", "chuva", "total"]
    else:
        headers = ["competencia", "unidade", "colaborador", "tipo", "data", "quantidade", "status", "motivo"]
    detail.append(headers)
    for row in rows:
        detail.append([money(value) if key in {"total", "media"} else value for key, value in row.items()])
    issue = wb.create_sheet("Inconsistências")
    issue.append(headers)
    for row in inconsistencies:
        issue.append([row.get(key, "") for key in headers])
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="116149")
            cell.alignment = Alignment(horizontal="center")
        for column in ws.columns:
            width = min(max(max(len(str(c.value or "")) for c in column) + 2, 12), 42)
            ws.column_dimensions[column[0].column_letter].width = width
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def create_pdf(report_type, rows, inconsistencies, metadata):
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=12 * mm, leftMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
    styles = getSampleStyleSheet()
    story = [Paragraph(metadata["title"], styles["Title"]), Paragraph(f"Emissor: {metadata['user']} | Emissão: {metadata['issued_at']}", styles["Normal"]), Paragraph(f"Filtros: {metadata['filters_text']}", styles["Normal"]), Spacer(1, 8)]
    if report_type == "monthly":
        header = ["Competência", "Colaborador", "Unidade", "Dias normais", "Dias chuva", "Valor normal", "Valor chuva", "Total"]
        data = [header] + [[r["competencia"], r["colaborador"], r["unidade"], r["dias_normais"], r["dias_chuva"], money(r["normal"]), money(r["chuva"]), money(r["total"])] for r in rows]
    elif report_type == "event_costs":
        header = ["Competência", "Unidade", "Evento", "Status", "Qtd. eventos", "Valor calculado", "Média/evento"]
        data = [header] + [[r["competencia"], r["unidade"], r["evento"], r["status"], r["quantidade"], money(r["total"]), money(r["media"])] for r in rows]
    elif report_type == "cost":
        header = ["Competência", "Unidade", "Total"]
        data = [header] + [[r["competencia"], r["unidade"], money(r["total"])] for r in rows]
    else:
        header = ["Competência", "Unidade", "Colaborador", "Tipo", "Data", "Qtd.", "Status"]
        data = [header] + [[r["competencia"], r["unidade"], r["colaborador"], r["tipo"], r["data"], r["quantidade"], r["status"]] for r in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#116149")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 8), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(table)
    story.append(Spacer(1, 8))
    story.append(Paragraph(f"Linhas emitidas: {len(rows)} | Inconsistências identificadas: {len(inconsistencies)}", styles["Normal"]))
    if inconsistencies:
        story.append(Paragraph("Atenção: existem eventos pendentes ou rejeitados. Este relatório não representa fechamento definitivo.", styles["Normal"]))
    if report_type == "event_costs":
        story.append(Paragraph("Valores calculados registrados no sistema; não comprovam pagamento efetuado. Eventos sem resultado de cálculo não entram no total.", styles["Normal"]))
    doc.build(story)
    return output.getvalue()


def response_file(content, content_type, filename):
    response = HttpResponse(content, content_type=content_type)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
