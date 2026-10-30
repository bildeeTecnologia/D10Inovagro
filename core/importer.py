import hashlib
import re
import unicodedata
from calendar import monthrange
from collections import Counter
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


PDF_MAX_PAGES = 200
PDF_MAX_TEXT_LENGTH = 2_000_000
PDF_EMPLOYEE_HEADER = re.compile(
    r"Empr\.\s*:\s*(?P<code>\d+)\s*(?P<name>.*?)\s*Situa[^:\s]*\s*:",
    re.IGNORECASE,
)
PDF_EMPLOYEE_TOTALS = re.compile(
    r"ND:\s*\d+\s+Proventos:\s*(?P<earnings>-?[\d.,]+)\s+"
    r"Descontos:\s*(?P<deductions>-?[\d.,]+).*?"
    r"L[^q\s]*quido:\s*(?P<net>-?[\d.,]+)",
    re.IGNORECASE,
)


def _normalized(value):
    value = unicodedata.normalize("NFKD", str(value or ""))
    return "".join(char for char in value if not unicodedata.combining(char)).lower().strip()


def _digits(value):
    return re.sub(r"\D", "", str(value or ""))


def _column(headers, options):
    normalized = {_normalized(header): index for index, header in enumerate(headers)}
    for option in options:
        if _normalized(option) in normalized:
            return normalized[_normalized(option)]
    for index, header in enumerate(headers):
        if any(_normalized(option) in _normalized(header) for option in options):
            return index
    return None


def sha256_file(path):
    digest = hashlib.sha256()
    with open(path, "rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _metadata(rows):
    text = " ".join(str(value) for row in rows[:5] for value in row if value not in (None, ""))
    result = {}
    for label in ("EMPRESA", "CNPJ", "COMPETÊNCIA", "COMPETENCIA"):
        for row in rows[:5]:
            for value in row:
                if isinstance(value, str) and label in value.upper():
                    result[label.lower().replace("ê", "e")] = value.strip()
    return result


def inspect_workbook(path, sheet_name=None):
    suffix = Path(path).suffix.lower()
    if suffix != ".xlsx":
        raise ValueError("Nesta etapa, a prévia aceita somente arquivos XLSX.")
    wb = load_workbook(path, read_only=True, data_only=False)
    selected = wb[sheet_name] if sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]
    first_rows = [list(row) for row in selected.iter_rows(min_row=1, max_row=min(selected.max_row, 12), values_only=True)]
    header_index = max(range(len(first_rows)), key=lambda index: sum(1 for value in first_rows[index] if isinstance(value, str) and value.strip()))
    headers = [str(value).strip() if value not in (None, "") else f"COLUNA_{index + 1}" for index, value in enumerate(first_rows[header_index])]
    preview = []
    for values in first_rows[header_index + 1:header_index + 6]:
        item = {}
        for index, header in enumerate(headers):
            value = values[index] if index < len(values) else ""
            if "CPF" in header.upper():
                value = "***.***.***-**" if value not in (None, "") else ""
            item[header] = "" if value is None else str(value)
        if any(item.values()):
            preview.append(item)
    errors = []
    if not any("FUNCION" in header.upper() for header in headers):
        errors.append("Não foi identificada uma coluna de funcionário.")
    if not any("UNIDADE" in header.upper() for header in headers):
        errors.append("Não foi identificada uma coluna de unidade.")
    return {"sheet_name": selected.title, "sheet_names": wb.sheetnames, "header_row": header_index + 1, "headers": headers, "preview": preview, "detected": _metadata(first_rows), "row_count": max(selected.max_row - header_index - 1, 0), "errors": errors}


def _pdf_amount(value):
    try:
        return Decimal(value.replace(".", "").replace(",", ".")).quantize(Decimal("0.01"))
    except (InvalidOperation, AttributeError) as exc:
        raise ValueError("O PDF contém um total monetário inválido.") from exc


def _pdf_metadata(text):
    company = re.search(r"Empresa:\s*\d+\s*-\s*(.*?)\s+P[aá]gina:", text, re.IGNORECASE)
    cnpj = re.search(r"CNPJ:\s*([\d./-]+)", text, re.IGNORECASE)
    competency = re.search(r"Compet[eê]ncia:\s*(\d{1,2})\s*/\s*(20\d{2})", text, re.IGNORECASE)
    return {
        "company_name": company.group(1).strip() if company else "",
        "cnpj": cnpj.group(1).strip() if cnpj else "",
        "month": int(competency.group(1)) if competency else None,
        "year": int(competency.group(2)) if competency else None,
    }


def _parse_payroll_pages(pages):
    if not pages:
        raise ValueError("O PDF não possui páginas.")
    if len(pages) > PDF_MAX_PAGES:
        raise ValueError(f"O PDF ultrapassa o limite de {PDF_MAX_PAGES} páginas.")
    if sum(len(text) for text in pages) > PDF_MAX_TEXT_LENGTH:
        raise ValueError("O texto extraído do PDF ultrapassa o limite permitido.")

    errors = []
    metadata_by_page = [_pdf_metadata(text) for text in pages]
    metadata = {}
    for key, default in (
        ("company_name", ""),
        ("cnpj", ""),
        ("month", None),
        ("year", None),
    ):
        metadata[key] = next(
            (item[key] for item in metadata_by_page if item.get(key)),
            default,
        )
    if not metadata.get("cnpj"):
        errors.append("CNPJ não identificado no PDF.")
    if not metadata.get("month") or not metadata.get("year"):
        errors.append("Competência não identificada no PDF.")
    cnpjs = {item["cnpj"] for item in metadata_by_page if item.get("cnpj")}
    if len(cnpjs) > 1:
        errors.append("O PDF apresenta mais de um CNPJ.")
    competencies = {
        (item["month"], item["year"])
        for item in metadata_by_page
        if item.get("month") and item.get("year")
    }
    if len(competencies) > 1:
        errors.append("O PDF apresenta mais de uma competência.")

    employees = []
    for page_number, text in enumerate(pages, start=1):
        headers = list(PDF_EMPLOYEE_HEADER.finditer(text))
        for index, header in enumerate(headers):
            block_end = headers[index + 1].start() if index + 1 < len(headers) else len(text)
            block = text[header.end():block_end]
            totals = PDF_EMPLOYEE_TOTALS.search(block)
            employee = {
                "row": len(employees) + 1,
                "page": page_number,
                "employee_code": header.group("code").strip(),
                "employee": header.group("name").strip(),
                "status": "invalid",
                "message": "",
            }
            if not totals:
                employee["message"] = "Totais do colaborador não identificados."
                errors.append("Há registro de colaborador sem totais legíveis.")
            else:
                earnings = _pdf_amount(totals.group("earnings"))
                deductions = _pdf_amount(totals.group("deductions"))
                net = _pdf_amount(totals.group("net"))
                employee.update({
                    "pdf_earnings": str(earnings),
                    "pdf_deductions": str(deductions),
                    "pdf_net": str(net),
                })
                if earnings - deductions != net:
                    employee["message"] = "Proventos menos descontos divergem do líquido informado."
                    errors.append("Há registro com totais individuais inconsistentes.")
                else:
                    employee.update(status="pending", message="Totais extraídos; aguardando conciliação.")
            employees.append(employee)

    if not employees:
        errors.append("Nenhum registro de colaborador foi identificado no PDF.")
    counts = Counter(item["employee_code"] for item in employees)
    if any(count > 1 for count in counts.values()):
        errors.append("Há códigos de colaborador repetidos no PDF.")
    for employee in employees:
        if counts[employee["employee_code"]] > 1:
            employee["status"] = "duplicate"
            employee["message"] = "Código de colaborador repetido no PDF; conciliação bloqueada."
    detected = {
        "company_name": metadata.get("company_name", ""),
        "cnpj": metadata.get("cnpj", ""),
        "competency": (
            f"{metadata['month']:02d}/{metadata['year']}"
            if metadata.get("month") and metadata.get("year")
            else ""
        ),
    }
    return {
        "metadata": metadata,
        "detected": detected,
        "employees": employees,
        "errors": errors,
        "page_count": len(pages),
    }


def inspect_payroll_pdf(path):
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        if not pdf.doc.is_extractable:
            raise ValueError("PDF protegido por senha não pode ser analisado.")
        parsed = _parse_payroll_pages([page.extract_text() or "" for page in pdf.pages])
    parsed["detected"]["source_format"] = "PDF"
    parsed["detected"]["page_count"] = parsed["page_count"]
    parsed["detected"]["row_count"] = len(parsed["employees"])
    parsed["preview"] = [
        {
            "Código": row["employee_code"],
            "Colaborador": row["employee"],
            "Proventos PDF": row.get("pdf_earnings", ""),
            "Descontos PDF": row.get("pdf_deductions", ""),
            "Líquido PDF": row.get("pdf_net", ""),
        }
        for row in parsed["employees"][:10]
    ]
    return parsed


def validate_payroll_pdf(inspection, tenant):
    from django.db.models import Q, Sum

    from .models import Collaborator, Company, Competency, EmploymentLink, PaymentCycle

    metadata = inspection["metadata"]
    cnpj = _digits(metadata.get("cnpj"))
    company = Company.objects.filter(tenant=tenant)
    company = next((item for item in company if _digits(item.cnpj) == cnpj), None) if cnpj else None
    competency = None
    if metadata.get("month") and metadata.get("year"):
        competency = Competency.objects.filter(
            tenant=tenant, month=metadata["month"], year=metadata["year"]
        ).first()

    errors = list(inspection["errors"])
    if not company:
        errors.append("Não foi possível vincular o CNPJ do PDF a uma empresa cadastrada.")
    if not competency:
        errors.append("A competência do PDF não está cadastrada para este cliente.")
    cycle = None
    if company and competency:
        cycle = PaymentCycle.objects.filter(
            tenant=tenant, competency=competency, cycle_type="monthly"
        ).first()
    aggregates = {}
    system_totals = None
    if cycle:
        aggregates = {
            item["collaborator_id"]: item
            for item in cycle.entries.values("collaborator_id").annotate(
                earnings=Sum("amount", filter=Q(rubric__category="earning")),
                deductions=Sum("amount", filter=Q(rubric__category="deduction")),
            )
        }
        all_totals = cycle.entries.aggregate(
            earnings=Sum("amount", filter=Q(rubric__category="earning")),
            deductions=Sum("amount", filter=Q(rubric__category="deduction")),
        )
        system_earnings = all_totals["earnings"] or Decimal("0.00")
        system_deductions = all_totals["deductions"] or Decimal("0.00")
        system_totals = {
            "earnings": str(system_earnings),
            "deductions": str(system_deductions),
            "net": str(system_earnings - system_deductions),
        }

    collaborators = {
        item.employee_code.strip(): item
        for item in Collaborator.objects.filter(tenant=tenant)
    }
    counts = Counter(row["employee_code"] for row in inspection["employees"])
    matched_count = 0
    discrepancy_count = 0
    unmatched_count = 0
    no_cycle_count = 0
    pdf_totals = {key: Decimal("0.00") for key in ("earnings", "deductions", "net")}
    for row in inspection["employees"]:
        if not all(key in row for key in ("pdf_earnings", "pdf_deductions", "pdf_net")):
            continue
        pdf_totals["earnings"] += Decimal(row["pdf_earnings"])
        pdf_totals["deductions"] += Decimal(row["pdf_deductions"])
        pdf_totals["net"] += Decimal(row["pdf_net"])
        if counts[row["employee_code"]] > 1 or row.get("status") in {"duplicate", "invalid"}:
            continue
        collaborator = None
        matched_links = EmploymentLink.objects.filter(
            collaborator__tenant=tenant,
            company=company,
            employee_code=row["employee_code"],
        ) if company and row["employee_code"] else EmploymentLink.objects.none()
        matched_collaborator_ids = list(matched_links.values_list("collaborator_id", flat=True).distinct())
        if len(matched_collaborator_ids) == 1:
            collaborator = Collaborator.objects.get(pk=matched_collaborator_ids[0])
        elif len(matched_collaborator_ids) > 1:
            row.update(status="unmatched", message="Matrícula ambígua: mais de um vínculo nesta empresa.")
            unmatched_count += 1
            continue
        if collaborator is None:
            collaborator = collaborators.get(row["employee_code"])
        if not collaborator:
            row.update(status="unmatched", message="Código interno não encontrado no cadastro.")
            unmatched_count += 1
            continue
        if company:
            employment_links = EmploymentLink.objects.filter(
                collaborator=collaborator, company=company
            )
            if competency:
                period_start = date(competency.year, competency.month, 1)
                period_end = date(
                    competency.year,
                    competency.month,
                    monthrange(competency.year, competency.month)[1],
                )
                employment_links = employment_links.filter(
                    admission_date__lte=period_end,
                ).filter(
                    Q(termination_date__isnull=True) | Q(termination_date__gte=period_start)
                )
            if not employment_links.exists():
                row.update(
                    status="company_mismatch",
                    message="Não há vínculo vigente na competência com a empresa do PDF.",
                )
                unmatched_count += 1
                continue
        if not cycle:
            row.update(status="no_cycle", message="Ciclo mensal interno ainda não foi gerado.")
            no_cycle_count += 1
            continue

        internal = aggregates.get(collaborator.id, {})
        internal_earnings = internal.get("earnings") or Decimal("0.00")
        internal_deductions = internal.get("deductions") or Decimal("0.00")
        internal_net = internal_earnings - internal_deductions
        row.update({
            "system_earnings": str(internal_earnings),
            "system_deductions": str(internal_deductions),
            "system_net": str(internal_net),
            "earnings_difference": str(Decimal(row["pdf_earnings"]) - internal_earnings),
            "deductions_difference": str(Decimal(row["pdf_deductions"]) - internal_deductions),
            "net_difference": str(Decimal(row["pdf_net"]) - internal_net),
        })
        has_difference = any(
            Decimal(row[field]) != 0
            for field in ("earnings_difference", "deductions_difference", "net_difference")
        )
        if has_difference:
            discrepancy_count += 1
        if _normalized(row["employee"]) != _normalized(collaborator.name):
            row.update(status="name_mismatch", message="Código corresponde, mas o nome diverge do cadastro.")
            matched_count += 1
        elif has_difference:
            row.update(status="discrepancy", message="Há diferença entre os totais do PDF e do sistema.")
            matched_count += 1
        else:
            row.update(status="reconciled", message="Totais do PDF e do sistema coincidem.")
            matched_count += 1

    summary = {
        "employee_count": len(inspection["employees"]),
        "matched_count": matched_count,
        "discrepancy_count": discrepancy_count,
        "unmatched_count": unmatched_count,
        "no_cycle_count": no_cycle_count,
        "duplicate_count": sum(count > 1 for count in counts.values()),
        "invalid_count": sum(row.get("status") == "invalid" for row in inspection["employees"]),
        "pdf_totals": {key: str(value) for key, value in pdf_totals.items()},
        "system_totals": system_totals,
        "total_differences": (
            {
                key: str(pdf_totals[key] - Decimal(system_totals[key]))
                for key in ("earnings", "deductions", "net")
            }
            if system_totals
            else None
        ),
    }
    if not cycle and company and competency:
        errors.append("Ciclo mensal interno não encontrado; valores não foram comparados.")
    detected = {
        **inspection["detected"],
        "matched_rows": matched_count,
        "summary": summary,
        "rubric_mapping": "pending",
    }
    return {
        "company": company,
        "competency": competency,
        "detected": detected,
        "validation_rows": inspection["employees"],
        "errors": errors,
    }


def validate_workbook(path, tenant, sheet_name=None, header_row=None):
    """Valida referências da planilha sem criar ou alterar registros de negócio."""
    from .models import Collaborator, Company, Competency, EmploymentLink, Unit

    wb = load_workbook(path, read_only=True, data_only=True)
    selected = wb[sheet_name] if sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = [list(row) for row in selected.iter_rows(values_only=True)]
    if not rows:
        return {"errors": ["A aba selecionada não possui linhas."], "company": None, "competency": None, "matched_rows": 0, "validation_rows": []}
    header_index = (header_row - 1) if header_row else max(range(min(len(rows), 12)), key=lambda index: sum(1 for value in rows[index] if isinstance(value, str) and value.strip()))
    headers = [str(value).strip() if value not in (None, "") else f"COLUNA_{index + 1}" for index, value in enumerate(rows[header_index])]
    name_index = _column(headers, ["funcionario", "colaborador", "empregado", "nome"])
    unit_index = _column(headers, ["unidade", "fazenda", "base"])
    code_index = _column(headers, ["matricula", "codigo de funcionario", "codigo funcionario", "codigo interno"])
    errors = []
    if name_index is None:
        errors.append("Não foi possível vincular a coluna de funcionário.")
    if unit_index is None:
        errors.append("Não foi possível vincular a coluna de unidade.")
    metadata = _metadata(rows[:8])
    cnpj = _digits(metadata.get("cnpj"))
    companies = list(Company.objects.filter(tenant=tenant))
    company = next((item for item in companies if cnpj and _digits(item.cnpj) == cnpj), None) if cnpj else (companies[0] if len(companies) == 1 else None)
    if not company:
        errors.append("Não foi possível identificar a empresa pelo CNPJ.")
    competency = None
    competency_text = metadata.get("competência") or metadata.get("competencia") or ""
    match = re.search(r"(0?[1-9]|1[0-2])\s*[/.-]\s*(20\d{2})", competency_text)
    if match:
        competency = Competency.objects.filter(tenant=tenant, month=int(match.group(1)), year=int(match.group(2))).first()
    if not competency:
        errors.append("Competência não identificada ou não cadastrada.")
    collaborators = list(Collaborator.objects.filter(tenant=tenant, active=True))
    units = list(Unit.objects.filter(tenant=tenant))
    matched_rows = 0
    validation_rows = []
    for row_number, values in enumerate(rows[header_index + 1:], start=header_index + 2):
        if not any(value not in (None, "") for value in values):
            continue
        name_value = values[name_index] if name_index is not None and name_index < len(values) else ""
        unit_value = values[unit_index] if unit_index is not None and unit_index < len(values) else ""
        code_value = str(values[code_index] or "").strip() if code_index is not None and code_index < len(values) else ""
        unit = next((item for item in units if _normalized(item.name) == _normalized(unit_value)), None)
        collaborator = None
        if code_value and company and unit:
            linked_ids = list(EmploymentLink.objects.filter(
                collaborator__tenant=tenant,
                company=company,
                unit=unit,
                employee_code__iexact=code_value,
            ).values_list("collaborator_id", flat=True).distinct())
            if len(linked_ids) == 1:
                collaborator = Collaborator.objects.get(pk=linked_ids[0])
            elif len(linked_ids) > 1:
                message = "Matrícula ambígua para esta empresa/unidade."
                errors.append(f"Linha {row_number}: {message}")
                validation_rows.append({"row": row_number, "employee": str(name_value or code_value), "unit": str(unit_value or ""), "status": "error", "message": message})
                continue
        if collaborator is None:
            lookup_value = code_value or name_value
            collaborator = next((item for item in collaborators if _normalized(item.name) == _normalized(lookup_value) or _normalized(item.employee_code) == _normalized(lookup_value)), None)
        if not collaborator:
            message = f"Funcionário não cadastrado ({name_value})."
            errors.append(f"Linha {row_number}: {message}")
            validation_rows.append({"row": row_number, "employee": str(name_value or ""), "unit": str(unit_value or ""), "status": "error", "message": message})
            continue
        if not unit:
            message = f"Unidade não cadastrada ({unit_value})."
            errors.append(f"Linha {row_number}: {message}")
            validation_rows.append({"row": row_number, "employee": str(name_value or ""), "unit": str(unit_value or ""), "status": "error", "message": message})
            continue
        if company and not EmploymentLink.objects.filter(collaborator=collaborator, company=company, unit=unit).exists():
            message = f"Vínculo de {collaborator.name} não corresponde à empresa/unidade informada."
            errors.append(f"Linha {row_number}: {message}")
            validation_rows.append({"row": row_number, "employee": collaborator.name, "unit": unit.name, "status": "error", "message": message})
            continue
        matched_rows += 1
        validation_rows.append({"row": row_number, "employee": collaborator.name, "unit": unit.name, "status": "matched", "message": "Correspondência encontrada."})
    return {"errors": errors, "company": company, "competency": competency, "matched_rows": matched_rows, "headers": headers, "validation_rows": validation_rows}
