import hashlib
import hmac
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal

from django.conf import settings
from django.db import IntegrityError, transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import Collaborator, Company, EmploymentLink, Unit

HEADERS = [
    "CNPJ da empresa",
    "Unidade",
    "Nome do colaborador",
    "CPF",
    "Matrícula no vínculo",
    "Cargo/função",
    "Data de admissão",
    "Data de desligamento",
]
MAX_FILE_SIZE = 15 * 1024 * 1024


def normalize_text(value):
    raw = unicodedata.normalize("NFKD", str(value or ""))
    return re.sub(r"[^a-z0-9]+", " ", "".join(c for c in raw if not unicodedata.combining(c)).lower()).strip()


def digits(value):
    if isinstance(value, (int, Decimal)):
        value = str(int(value))
    elif isinstance(value, float) and value.is_integer():
        value = str(int(value))
    elif isinstance(value, str) and re.fullmatch(r"\d+\.0+", value.strip()):
        value = value.strip().split(".", 1)[0]
    return re.sub(r"\D", "", str(value or ""))


def valid_cpf(value):
    cpf = digits(value).zfill(11)
    if len(cpf) != 11 or len(set(cpf)) == 1:
        return False
    for length, weight in ((9, 10), (10, 11)):
        total = sum(int(cpf[index]) * (weight - index) for index in range(length))
        check = (total * 10) % 11
        if check == 10:
            check = 0
        if check != int(cpf[length]):
            return False
    return True


def cpf_fingerprint(value):
    cpf = digits(value).zfill(11)
    key = settings.COLLABORATOR_ID_KEY.encode("utf-8")
    return hmac.new(key, b"agro-collaborator-cpf-v1:" + cpf.encode("ascii"), hashlib.sha256).hexdigest()


def parse_date(value, required=False):
    if value in (None, ""):
        if required:
            raise ValueError("Data de admissão obrigatória.")
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError("Data inválida; use dd/mm/aaaa.")


def create_roster_template():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Vinculos"
    sheet.append(HEADERS)
    sheet.append(["00.000.000/0000-00", "Exemplo: Unidade 1", "Exemplo Colaborador", "00000000000", "M-001", "Cargo", "01/01/2026", ""])
    sheet.append([])
    sheet.append(["INSTRUÇÕES"])
    instructions = [
        "Uma linha por vínculo pessoa–empresa–unidade. Para a mesma pessoa em duas unidades, repita o CPF em duas linhas, uma por vínculo.",
        "CPF é usado somente para reconhecer a mesma pessoa. O sistema guarda uma impressão digital protegida, não o CPF aberto nem a forma mascarada.",
        "Matrícula pertence ao vínculo e pode ser diferente entre empresas/unidades. Se não existir, deixe em branco; não use o número sequencial da linha.",
        "CNPJ deve corresponder ao cadastro da empresa no sistema; a unidade é cadastrada dentro dessa empresa.",
        "Data de admissão é obrigatória; data de desligamento é opcional. Não preencha salário, CPF de dependentes, dados bancários ou descontos neste arquivo.",
        "Use a aba Vinculos como fonte de dados. A linha de exemplo deve ser removida antes do envio.",
        "O carregamento começa em prévia; nada é gravado até a confirmação explícita do lote.",
    ]
    for instruction in instructions:
        sheet.append([instruction])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:H2"
    table = Table(displayName="CargaVinculos", ref="A1:H2")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)
    widths = [24, 30, 36, 18, 22, 28, 20, 22]
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64+i)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for cell in sheet[2]:
        cell.fill = PatternFill("solid", fgColor="FFF2CC")
    sheet.row_dimensions[1].height = 32
    sheet.sheet_view.showGridLines = False
    return workbook


def _header_indexes(headers):
    aliases = {
        "cnpj": ("cnpj da empresa", "cnpj"),
        "unit": ("unidade", "nome da unidade", "fazenda", "base"),
        "name": ("nome do colaborador", "colaborador", "funcionario", "empregado", "nome"),
        "cpf": ("cpf", "cpf do colaborador"),
        "code": ("matricula no vinculo", "matricula", "codigo do vinculo", "codigo funcionario"),
        "job": ("cargo funcao", "cargo", "funcao"),
        "admission": ("data de admissao", "admissao"),
        "termination": ("data de desligamento", "desligamento", "data de demissao"),
    }
    normalized = [normalize_text(header) for header in headers]
    found = {}
    for key, options in aliases.items():
        for index, header in enumerate(normalized):
            if header in options:
                found[key] = index
                break
    return found


def prepare_roster_import(file_obj, tenant):
    if getattr(file_obj, "size", 0) > MAX_FILE_SIZE:
        raise ValueError("O arquivo não pode ultrapassar 15 MB.")
    if not str(getattr(file_obj, "name", "")).lower().endswith(".xlsx"):
        raise ValueError("Envie um arquivo XLSX.")
    file_obj.seek(0)
    raw_bytes = file_obj.read()
    if len(raw_bytes) > MAX_FILE_SIZE:
        raise ValueError("O arquivo não pode ultrapassar 15 MB.")
    digest = hashlib.sha256(raw_bytes).hexdigest()
    import io
    workbook = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    if not workbook.worksheets:
        raise ValueError("O arquivo não possui abas.")
    sheet = workbook["Vinculos"] if "Vinculos" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    header_row = next((i for i, row in enumerate(rows[:12]) if "cpf" in _header_indexes(row) and "unit" in _header_indexes(row)), None)
    if header_row is None:
        workbook.close()
        raise ValueError("Não encontrei cabeçalho com CPF e Unidade. Use o modelo disponibilizado.")
    indexes = _header_indexes(rows[header_row])
    missing = [key for key in ("cnpj", "unit", "name", "cpf", "admission") if key not in indexes]
    if missing:
        workbook.close()
        raise ValueError("Colunas obrigatórias ausentes: " + ", ".join(missing))
    companies = {re.sub(r"\D", "", company.cnpj): company for company in Company.objects.filter(tenant=tenant)}
    prepared, errors, warnings = [], [], []
    seen = {}
    person_names = {}
    for row_number, values in enumerate(rows[header_row + 1:], start=header_row + 2):
        if not any(value not in (None, "") for value in values):
            continue
        def val(key):
            index = indexes.get(key)
            return values[index] if index is not None and index < len(values) else None
        name = str(val("name") or "").strip()
        cnpj_digits = digits(val("cnpj"))
        unit_name = str(val("unit") or "").strip()
        cpf_raw = val("cpf")
        cpf_digits = digits(cpf_raw).zfill(11)
        row_errors = []
        if not name:
            row_errors.append("Nome ausente.")
        if len(cnpj_digits) != 14 or cnpj_digits not in companies:
            row_errors.append("CNPJ inválido ou não cadastrado para este cliente.")
        if not unit_name:
            row_errors.append("Unidade ausente.")
        if not valid_cpf(cpf_digits):
            row_errors.append("CPF inválido; corrija na fonte antes de carregar.")
        try:
            admission = parse_date(val("admission"), required=True)
            termination = parse_date(val("termination"))
            if termination and termination < admission:
                raise ValueError("Desligamento anterior à admissão.")
        except ValueError as exc:
            admission = None
            termination = None
            row_errors.append(str(exc))
        if row_errors:
            errors.extend([f"Linha {row_number}: {message}" for message in row_errors])
            prepared.append({"row": row_number, "name": name, "company": str(val("cnpj") or ""), "unit": unit_name, "status": "error", "message": "; ".join(row_errors)})
            continue
        company = companies[cnpj_digits]
        fingerprint = cpf_fingerprint(cpf_digits)
        code = str(val("code") or "").strip()
        job = str(val("job") or "").strip()
        key = (fingerprint, company.pk, normalize_text(unit_name), admission.isoformat(), termination.isoformat() if termination else "")
        person_names.setdefault(fingerprint, [])
        if name not in person_names[fingerprint]:
            person_names[fingerprint].append(name)
        if key in seen:
            existing = seen[key]
            existing["source_rows"].append(row_number)
            if existing["name"] != name:
                existing["name_variants"].append(name)
            if code and not existing.get("employee_code"):
                existing["employee_code"] = code
            if job and not existing.get("job_title"):
                existing["job_title"] = job
            if admission and admission < date.fromisoformat(existing["admission_date"]):
                existing["admission_date"] = admission.isoformat()
            if termination and (not existing.get("termination_date") or termination > date.fromisoformat(existing["termination_date"])):
                existing["termination_date"] = termination.isoformat()
            continue
        existing_person = Collaborator.objects.filter(tenant=tenant, cpf_fingerprint=fingerprint).first()
        existing_unit = next((item for item in Unit.objects.filter(tenant=tenant, company=company) if normalize_text(item.name) == normalize_text(unit_name)), None)
        link_exists = bool(existing_person and existing_unit and EmploymentLink.objects.filter(
            collaborator=existing_person,
            company=company,
            unit=existing_unit,
            admission_date=admission,
            termination_date=termination,
        ).exists())
        item = {
            "row": row_number,
            "source_rows": [row_number],
            "cpf_fingerprint": fingerprint,
            "name": name,
            "name_variants": [],
            "company_id": company.pk,
            "company_name": company.legal_name,
            "company_cnpj": company.cnpj,
            "unit": unit_name,
            "employee_code": code,
            "job_title": job,
            "admission_date": admission.isoformat(),
            "termination_date": termination.isoformat() if termination else "",
            "status": "existing_link" if link_exists else "ready",
            "message": "Vínculo já cadastrado; será ignorado." if link_exists else "Pronto para importar.",
        }
        seen[key] = item
        prepared.append(item)
    workbook.close()
    for item in prepared:
        if item.get("status") == "error":
            continue
        variants = person_names.get(item["cpf_fingerprint"], [])
        if len(variants) > 1:
            item["name"] = variants[0]
            item["status"] = "review_name"
            item["message"] = "Mesmo CPF com grafias diferentes do nome; será usada a primeira grafia da planilha. Confira após a carga."
            item["name_variants"] = sorted(variants)
            warnings.append(f"Pessoa com variação de nome em {len(variants)} grafias; CPF foi usado para consolidar.")
    # Duplicate rows have been collapsed into one link. Their line numbers remain traceable.
    valid_rows = [item for item in prepared if item.get("status") != "error"]
    return {
        "sha256": digest,
        "sheet_name": sheet.title,
        "row_count": len(rows) - header_row - 1,
        "rows": valid_rows,
        "errors": errors,
        "warnings": sorted(set(warnings)),
        "ready_count": sum(item["status"] in {"ready", "review_name"} for item in valid_rows),
        "existing_count": sum(item["status"] == "existing_link" for item in valid_rows),
        "error_count": sum(item.get("status") == "error" for item in prepared),
        "unit_count": len({(item.get("company_id"), normalize_text(item.get("unit"))) for item in valid_rows}),
        "person_count": len({item.get("cpf_fingerprint") for item in valid_rows}),
    }


def apply_roster_import(batch, user):
    if batch.status != "preview":
        raise ValueError("Este lote já foi confirmado ou rejeitado.")
    if batch.errors:
        raise ValueError("Corrija os erros da prévia antes de confirmar a carga.")
    tenant = batch.tenant
    created_people = created_links = created_units = skipped_links = 0
    with transaction.atomic():
        for catalog_item in batch.detected.get("unit_catalog", []):
            company = Company.objects.get(pk=catalog_item["company_id"], tenant=tenant)
            wanted_unit = catalog_item["unit"].strip()
            unit = next((u for u in Unit.objects.filter(tenant=tenant, company=company) if normalize_text(u.name) == normalize_text(wanted_unit)), None)
            if unit is None:
                Unit.objects.create(tenant=tenant, company=company, name=wanted_unit)
                created_units += 1
        for item in batch.validation_rows:
            if item.get("status") == "error":
                raise ValueError(f"A linha {item.get('row')} contém erro e não pode ser importada.")
            company = Company.objects.get(pk=item["company_id"], tenant=tenant)
            if re.sub(r"\D", "", company.cnpj) != re.sub(r"\D", "", item["company_cnpj"]):
                raise ValueError("O CNPJ da empresa mudou desde a prévia; gere outra prévia.")
            wanted_unit = item["unit"].strip()
            unit = next((u for u in Unit.objects.filter(tenant=tenant, company=company) if normalize_text(u.name) == normalize_text(wanted_unit)), None)
            if unit is None:
                unit = Unit.objects.create(tenant=tenant, company=company, name=wanted_unit)
                created_units += 1
            collaborator = Collaborator.objects.filter(tenant=tenant, cpf_fingerprint=item["cpf_fingerprint"]).first()
            if collaborator is None:
                collaborator = Collaborator.objects.create(
                    tenant=tenant,
                    name=item["name"],
                    employee_code="",
                    cpf_fingerprint=item["cpf_fingerprint"],
                    active=True,
                )
                created_people += 1
            elif normalize_text(collaborator.name) != normalize_text(item["name"]):
                # CPF is the identity key, but an existing name is never silently overwritten.
                raise ValueError(f"O nome cadastrado diverge do CPF na linha {item.get('row')}; revise o cadastro antes de confirmar.")
            admission = date.fromisoformat(item["admission_date"])
            termination = date.fromisoformat(item["termination_date"]) if item.get("termination_date") else None
            link = EmploymentLink.objects.filter(
                collaborator=collaborator,
                company=company,
                unit=unit,
                admission_date=admission,
                termination_date=termination,
            ).first()
            if link:
                skipped_links += 1
                continue
            EmploymentLink.objects.create(
                collaborator=collaborator,
                company=company,
                unit=unit,
                employee_code=item.get("employee_code", ""),
                job_title=item.get("job_title", "") or "A confirmar",
                admission_date=admission,
                termination_date=termination,
            )
            created_links += 1
        batch.status = "confirmed"
        batch.confirmed_by = user
        batch.confirmed_at = timezone.now()
        batch.detected = {**batch.detected, "created_people": created_people, "created_units": created_units, "created_links": created_links, "skipped_links": skipped_links}
        # Keep only aggregate results after confirmation; drop fingerprints and names from the import staging batch.
        batch.validation_rows = []
        batch.preview = []
        batch.save(update_fields=["status", "confirmed_by", "confirmed_at", "detected", "validation_rows", "preview", "updated_at"])
    return {"created_people": created_people, "created_units": created_units, "created_links": created_links, "skipped_links": skipped_links}
